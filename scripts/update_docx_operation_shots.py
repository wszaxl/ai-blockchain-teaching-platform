# -*- coding: utf-8 -*-
"""Replace section 4.3 (操作试验图) images in the proposal docx with real platform
operation screenshots (template-style): topology/chain-management shot + explorer
block/transaction-detail shot.

Usage: python scripts/update_docx_operation_shots.py <img1> <cap1> <img2> <cap2>
"""
import os
import sys
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = "E:/HuaweiMoveData/Users/HUAWEI/Desktop/基于AL+区域协同教研"
DOCX = os.path.join(ROOT, "docs", "教研平台", "基于AI+区块链的区域协同教研平台应用创新方案.docx")
WPD = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"

def has_img(p):
    return bool(p._element.findall(".//%sinline" % WPD)) or \
           bool(p._element.findall(".//%sanchor" % WPD))

def main(img1, cap1, img2, cap2):
    d = docx.Document(DOCX)
    paras = d.paragraphs
    i43 = None
    for i, p in enumerate(paras):
        if p.style.name == "Heading 2" and p.text.strip().startswith("4.3"):
            i43 = i
            break
    assert i43 is not None, "4.3 heading not found"
    iend = None
    for i in range(i43 + 1, len(paras)):
        if paras[i].style.name in ("Heading 1", "Heading 2"):
            iend = i
            break
    assert iend is not None

    # remove old images + their captions inside 4.3
    removed = 0
    for i in range(i43 + 1, iend):
        p = paras[i]
        if has_img(p) or p.text.strip().startswith("图 4-"):
            p._element.getparent().remove(p._element)
            removed += 1
    print("removed old paras:", removed)

    # anchor = first remaining body paragraph in 4.3, else the next heading
    paras = d.paragraphs
    i43 = next(i for i, p in enumerate(paras)
               if p.style.name == "Heading 2" and p.text.strip().startswith("4.3"))
    iend = next(i for i in range(i43 + 1, len(paras))
                if paras[i].style.name in ("Heading 1", "Heading 2"))
    anchor = None
    for i in range(i43 + 1, iend):
        if paras[i].text.strip():
            anchor = paras[i]
            break
    if anchor is None:
        anchor = paras[iend]

    for img, cap in ((img1, cap1), (img2, cap2)):
        pic = anchor.insert_paragraph_before()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(img, width=Inches(6.2))
        c = anchor.insert_paragraph_before()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap)
        r.font.size = Pt(10)

    d.save(DOCX)
    print("saved:", DOCX)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
