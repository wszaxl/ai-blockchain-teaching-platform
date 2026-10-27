# -*- coding: utf-8 -*-
"""在第 5 章追加开发关键节点小节（5.1-5.4），在 6.2 追加可视化大屏图。"""
import os
import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(WS, 'docs', '教研平台', '基于AI+区块链的区域协同教研平台应用创新方案.docx')
ASSETS = os.path.join(WS, 'docs', '教研平台', 'assets')

d = docx.Document(DOC)


def find(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise SystemExit('anchor not found: ' + prefix)


def insert_after(prev, style=None, text=None, image=None, width=6.2):
    new_p = docx.oxml.OxmlElement('w:p')
    prev._p.addnext(new_p)
    np = Paragraph(new_p, prev._parent)
    if style:
        np.style = d.styles[style]
    if text:
        np.add_run(text)
    if image:
        np.add_run().add_picture(os.path.join(ASSETS, image), width=Inches(width))
    return np


anchor5 = find('图 5-2')
seq = [
    ('Heading 2', '5.1 开发环境、工具链与项目整体结构', None),
    ('Normal', '项目按「合约 — 工具链 — 网关 — 文档/可视化」四层组织：contracts/ 存放 Solidity 源码，tools/ 为纯 JDK 链上工具链，'
               'build/ 保存编译器与运行产物，可视化/ 与 docs/ 消费同一批真实回执。私钥 keys/access.key 仅存于演示机，提交压缩包时删除且内容全程打码。整体结构见图 5-3。', None),
    (None, None, 'fig5-3_项目整体结构.png'),
    ('Normal', '图 5-3　项目整体结构与角色标注（真实目录树）', None),
    ('Heading 2', '5.2 编译与部署关键节点（蚂蚁链 BaaS REST 网关）', None),
    ('Normal', '合约经蚂蚁链官方定制编译器 alipay-solc-0.4.24 编译（并与标准 solc 0.4.26 交叉核验字节码一致）；部署通过蚂蚁链开放联盟链开发平台的 '
               'BaaS REST 网关 rest.baas.alipay.com 完成 SHA256withRSA 签名与交易发送，链账户为 KMS 托管账户（名称打码）。编译与部署回执见图 5-4。', None),
    (None, None, 'fig5-4_编译与部署终端.png'),
    ('Normal', '图 5-4　开发关键节点：合约编译与部署上链（真实终端输出，链账户名称已打码）', None),
    ('Heading 2', '5.3 业务上链与回执核验关键节点', None),
    ('Normal', '成果确权与教研活动两笔交易上链后，使用 QUERYRECEIPT 拉取回执，并将 logData 的 base64 解码即得链上存证原文 miniappinfo，'
               '作者、学校、学分等字段与第 6.2 节凭证表逐项一致；两笔哈希亦可在 ChainInsight 浏览器复核。见图 5-5。', None),
    (None, None, 'fig5-5_上链回执核验终端.png'),
    ('Normal', '图 5-5　开发关键节点：上链回执核验与存证原文解码（实时查链输出）', None),
    ('Heading 2', '5.4 链上只读查询关键节点：学分折算现场验证', None),
    ('Normal', 'getCreditFor(string) 为只读调用（CALLCONTRACTBIZ + isLocalTransaction=true），不打包上链、不消耗 gas。'
               '现场查询教师「陈以宁」返回 180，与 90 分钟 × 2 的合约折算口径一致，验证链上业务闭环。见图 5-6。', None),
    (None, None, 'fig5-6_学分查询终端.png'),
    ('Normal', '图 5-6　开发关键节点：getCreditFor 学分现场查询（只读、不耗 gas）', None),
]
prev = anchor5
for style, text, image in seq:
    prev = insert_after(prev, style, text, image)

anchor6 = find('以上交易均发生在蚂蚁链开放联盟链')
prev = anchor6
for style, text, image in [
    ('Normal', '上述凭证同时汇入本地链上交易可视化大屏（可视化/visualize_server.py，http://127.0.0.1:8765/）：页面卡片直接读取真实回执，'
               '并支持粘贴任意 64 位交易哈希经本地 RestChainTool 现场查链；另已发布静态只读版供在线查看。见图 6-1。', None),
    (None, None, 'fig6-1_可视化大屏.png'),
    ('Normal', '图 6-1　链上交易可视化大屏：miniappinfo 业务记录实时展示（真实页面截图）', None),
]:
    prev = insert_after(prev, style, text, image)

d.save(DOC)
print('saved; paragraphs =', len(d.paragraphs))
