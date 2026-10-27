# -*- coding: utf-8 -*-
"""Replace the plain-text smart-contract code block in section 5 of the proposal
docx with the dark-themed core-code screenshot images (fig4_代码核心截图_1/2.png)."""
import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = "E:/HuaweiMoveData/Users/HUAWEI/Desktop/基于AL+区域协同教研"
DOCX = os.path.join(ROOT, "docs", "教研平台", "基于AI+区块链的区域协同教研平台应用创新方案.docx")
BAK = DOCX + ".bak"
A1 = os.path.join(ROOT, "docs", "教研平台", "assets", "fig4_代码核心截图_1.png")
A2 = os.path.join(ROOT, "docs", "教研平台", "assets", "fig4_代码核心截图_2.png")

if not os.path.exists(BAK):
    shutil.copy2(DOCX, BAK)
    print("backup ->", BAK)

d = docx.Document(DOCX)
paras = d.paragraphs

# locate section-5 Heading 1 and the following Heading 1 (section 6)
i5 = None
for i, p in enumerate(paras):
    if (p.style.name == "Heading 1") and p.text.strip().startswith("5"):
        i5 = i
        break
assert i5 is not None, "section 5 heading not found"
i6 = None
for i in range(i5 + 1, len(paras)):
    if paras[i].style.name == "Heading 1":
        i6 = i
        break
assert i6 is not None, "section 6 heading not found"

intro = paras[i5 + 1]
code_paras = paras[i5 + 2:i6]
next_heading = paras[i6]
print("section5 idx", i5, "intro:", intro.text[:30], "code paras:", len(code_paras))

# remove the old plain-text code paragraphs
for p in code_paras:
    p._element.getparent().remove(p._element)

def add_image_para(img, caption):
    pic = next_heading.insert_paragraph_before()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(img, width=Inches(6.2))
    cap = next_heading.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.size = Pt(10)
    return pic, cap

add_image_para(A1, "图 5-1 教研合约链上数据结构（contracts/TeachingTrace.sol 核心代码截图）")
add_image_para(A2, "图 5-2 冻结 ABI 三入口：确权存证 / 活动留痕与学分折算 / 学分查询（核心代码截图）")

d.save(DOCX)
print("saved:", DOCX)
