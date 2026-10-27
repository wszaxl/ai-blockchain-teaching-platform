# -*- coding: utf-8 -*-
"""生成《基于AI+区块链的区域协同教研平台》方案 PPT（提纲版，可在 PPT/WPS 中继续润色）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(WS, 'docs', '教研平台', '教研平台-方案汇报.pptx')
TA = os.path.join(WS, 'docs', '教研平台')

ACCENT = RGBColor(0x0F, 0x4C, 0x81)
GRAY = RGBColor(0x6B, 0x72, 0x80)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_slide(title, bullets):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = ACCENT
    body = s.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11.5), Inches(5.4))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = para.add_run(); r.text = b
        r.font.size = Pt(17 if len(b) < 60 else 15)
        para.space_after = Pt(10)
    return s

add_slide('基于 AI+区块链的区域协同教研平台\n——技术创新组方案汇报', 
          ['院校：＿＿＿＿　 团队：＿＿＿＿', '区块链底座：蚂蚁链开放联盟链（a00e36c5，已真实上链）',
           'AI 层：资源推荐 / 会议纪要 / 教学分析 / 教师画像（模拟+预留）'])
add_slide('01 业务痛点', 
          ['① 跨校资源孤岛 + 版权归属不明，成果易被盗用',
           '② 过程记录靠人工签到/纸质台账，造假与代签难杜绝，贡献无法量化',
           '③ 各校学分口径不一，跨校学分互认缺失',
           '④ 资料分散、缺智能分析，选题与评估靠经验'])
add_slide('02 解决方案：双链架构', 
          ['AI 层提效：精准推荐、自动纪要、数据诊断、动态画像',
           '区块链层可信：成果确权、过程留痕、学分互认、审计追溯',
           'AI 处理结果与原始数据哈希上链 → 「AI 上链即审计」'])
add_slide('03 区块链应用点与网络', 
          ['应用点：成果确权 / 参与留痕与贡献统计 / 跨校学分互认 / 全流程审计',
           '网络：开放联盟链 a00e36c5 · 租户 LXHOXRFH · KMS 托管账户「陈以宁」',
           '（拓扑图/业务流程图/上链数据分析图见技术文档图 1~3）'])
add_slide('04 上链数据设计', 
          ['教研成果 Artifact：编号/标题/作者/学校/SHA-256 文件哈希',
           '教研活动 Activity：活动号/教师/学校/贡献分钟（×2 = 学分）',
           '展示格式：miniappinfo JSON，浏览器与小程序可直接渲染'])
add_slide('05 智能合约 TeachingTrace.sol', 
          ['registerArtifact(5×string)：成果确权存证',
           'recordActivity(4×string)：活动留痕 + 学分自动折算累加',
           'getCreditFor(string)：教师累计学分（本地只读）',
           '注释与代码逐段严格对应；接口文档独立成册'])
add_slide('06 真实链上凭证（SDK 调用）', 
          ['合约部署 ceeaba86…bd956b9（区块 183416200 系）',
           '成果确权 48ba4912…752de1bb；活动记录 7b844056…30fdb3cb',
           '学分查询 getCreditFor("陈以宁") = 180（90 分钟 × 2，本地只读）',
           '工具：tools/TeachFlowTool.java 四子命令，全程留档'])
add_slide('07 提交资料与视频计划', 
          ['PDF 技术文档：已生成（教研平台-技术文档.pdf，覆盖赛题全部检查点）',
           '视频三件套：搭建过程/合约运行/SDK 运行，分镜脚本+口播稿已备（07-过程实录）',
           '搭建口径：平台提供开放联盟链，以控制台配置+合约部署实录替代，理由书面化'])

prs.save(OUT)
print('PPT 已生成:', OUT, os.path.getsize(OUT), '字节，', len(prs.slides.__iter__.__self__._sldIdLst), '页')